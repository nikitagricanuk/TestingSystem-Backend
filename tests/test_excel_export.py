from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.export.excel_export import UnknownColumnError, build_workbook


class TestBuildWorkbook:
    def test_writes_header_and_rows_in_requested_order(self):
        rows = [{"name": "Ivan", "score": 90}, {"name": "Petr", "score": 80}]
        headers = {"name": "Имя", "score": "Балл"}

        content = build_workbook(rows, columns=["score", "name"], headers=headers)

        workbook = load_workbook(BytesIO(content))
        sheet = workbook.active
        assert [c.value for c in sheet[1]] == ["Балл", "Имя"]
        assert [c.value for c in sheet[2]] == [90, "Ivan"]
        assert [c.value for c in sheet[3]] == [80, "Petr"]

    def test_missing_row_value_becomes_blank_cell(self):
        rows = [{"name": "Ivan"}]
        headers = {"name": "Имя", "score": "Балл"}

        content = build_workbook(rows, columns=["name", "score"], headers=headers)

        sheet = load_workbook(BytesIO(content)).active
        # openpyxl reads an empty-string cell back as None (blank), not "".
        assert sheet[2][1].value is None

    def test_unknown_column_raises(self):
        with pytest.raises(UnknownColumnError):
            build_workbook([], columns=["bogus"], headers={"name": "Имя"})

    def test_sheet_title_is_truncated_to_excel_limit(self):
        content = build_workbook([], columns=["name"], headers={"name": "Имя"}, sheet_title="x" * 40)
        sheet = load_workbook(BytesIO(content)).active
        assert len(sheet.title) == 31
